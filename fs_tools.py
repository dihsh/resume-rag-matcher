"""
fs_tools.py - Document Parsing and File Operations for Resume RAG

Handles robust file I/O and text extraction across PDF, DOCX, and TXT formats
with metadata tracking and error resilience.
"""

from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    import docx
except ImportError:
    docx = None


def read_file(filepath: str) -> Dict[str, Any]:
    """
    Read document file (.pdf, .docx, .txt) and extract text content + metadata.
    """
    path = Path(filepath).resolve()

    if not path.exists():
        return {
            "status": "error",
            "filepath": str(path),
            "content": "",
            "metadata": {},
            "error": f"File not found: {path}",
        }

    if not path.is_file():
        return {
            "status": "error",
            "filepath": str(path),
            "content": "",
            "metadata": {},
            "error": f"Path is not a regular file: {path}",
        }

    stat = path.stat()
    ext = path.suffix.lower()
    metadata: Dict[str, Any] = {
        "file_name": path.name,
        "extension": ext,
        "size_bytes": stat.st_size,
        "modified_date": datetime.fromtimestamp(stat.st_mtime).isoformat(),
        "page_count": None,
    }

    try:
        content = ""
        if ext == ".txt":
            try:
                content = path.read_text(encoding="utf-8")
            except UnicodeDecodeError:
                content = path.read_text(encoding="latin-1")

        elif ext == ".pdf":
            if PdfReader is None:
                raise ImportError("pypdf is required to read PDF documents.")
            reader = PdfReader(str(path))
            metadata["page_count"] = len(reader.pages)
            pages_text = []
            for page in reader.pages:
                text = page.extract_text() or ""
                pages_text.append(text)
            content = "\n\n".join(pages_text)

        elif ext == ".docx":
            if docx is None:
                raise ImportError("python-docx is required to read DOCX documents.")
            doc = docx.Document(str(path))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            content = "\n".join(paragraphs)

        else:
            return {
                "status": "error",
                "filepath": str(path),
                "content": "",
                "metadata": metadata,
                "error": f"Unsupported extension: {ext}",
            }

        return {
            "status": "success",
            "filepath": str(path),
            "content": content.strip(),
            "metadata": metadata,
            "error": None,
        }

    except Exception as exc:
        return {
            "status": "error",
            "filepath": str(path),
            "content": "",
            "metadata": metadata,
            "error": f"Error reading {path.name}: {str(exc)}",
        }


def list_files(directory: str, extension: Optional[str] = None) -> List[Dict[str, Any]]:
    """
    List files in a directory with optional extension filtering.
    """
    dir_path = Path(directory).resolve()
    if not dir_path.exists() or not dir_path.is_dir():
        return []

    target_ext = None
    if extension:
        target_ext = extension.strip().lower()
        if not target_ext.startswith("."):
            target_ext = f".{target_ext}"

    results = []
    try:
        for entry in dir_path.iterdir():
            if entry.is_file():
                ext = entry.suffix.lower()
                if target_ext is None or ext == target_ext:
                    stat = entry.stat()
                    results.append({
                        "name": entry.name,
                        "path": str(entry.resolve()),
                        "size_bytes": stat.st_size,
                        "modified_date": datetime.fromtimestamp(stat.st_mtime).isoformat(),
                        "extension": ext,
                    })
    except Exception:
        return []

    return sorted(results, key=lambda x: x["name"].lower())


def write_file(filepath: str, content: str) -> Dict[str, Any]:
    """
    Write text content to file, creating directories if needed.
    """
    path = Path(filepath).resolve()
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        data = content.encode("utf-8")
        path.write_bytes(data)
        return {
            "status": "success",
            "filepath": str(path),
            "bytes_written": len(data),
            "message": f"Wrote {len(data)} bytes to {path.name}",
        }
    except Exception as exc:
        return {
            "status": "error",
            "filepath": str(path),
            "bytes_written": 0,
            "message": f"Write failed: {str(exc)}",
        }
