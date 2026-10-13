"""
generate_dataset.py - Generates 32 diverse resumes (PDF, DOCX, TXT) and 6 job descriptions.
"""

import json
from pathlib import Path
import docx
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

RESUMES_DIR = Path("resumes")
JDS_DIR = Path("job_descriptions")
DATA_DIR = Path("data")

RESUMES_DIR.mkdir(parents=True, exist_ok=True)
JDS_DIR.mkdir(parents=True, exist_ok=True)


def build_pdf(filepath: Path, name: str, contact: str, summary: str, skills: list, experience: list, education: str):
    doc = SimpleDocTemplate(str(filepath), pagesize=letter, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    t_style = ParagraphStyle('T', parent=styles['Heading1'], fontSize=16, leading=20, textColor=colors.HexColor('#1E3A8A'), spaceAfter=2)
    c_style = ParagraphStyle('C', parent=styles['Normal'], fontSize=9, textColor=colors.HexColor('#4B5563'), spaceAfter=6)
    h2_style = ParagraphStyle('H2', parent=styles['Heading2'], fontSize=11, leading=14, textColor=colors.HexColor('#111827'), spaceBefore=6, spaceAfter=3)
    b_style = ParagraphStyle('B', parent=styles['Normal'], fontSize=9, leading=12, textColor=colors.HexColor('#1F2937'), spaceAfter=2)

    story = [
        Paragraph(name, t_style),
        Paragraph(contact, c_style),
        Paragraph("Professional Summary", h2_style),
        Paragraph(summary, b_style),
        Paragraph("Technical Skills", h2_style),
        Paragraph(", ".join(skills), b_style),
        Paragraph("Work Experience", h2_style),
    ]
    for role, comp, dates, desc in experience:
        story.append(Paragraph(f"<b>{role}</b> &mdash; <i>{comp} ({dates})</i>", b_style))
        story.append(Paragraph(desc, b_style))
        story.append(Spacer(1, 2))
    story.append(Paragraph("Education", h2_style))
    story.append(Paragraph(education, b_style))
    doc.build(story)


def build_docx(filepath: Path, name: str, contact: str, summary: str, skills: list, experience: list, education: str):
    doc = docx.Document()
    doc.add_heading(name, level=0)
    p_c = doc.add_paragraph(contact)
    p_c.style.font.color.rgb = docx.shared.RGBColor(100, 100, 100)
    doc.add_heading("Professional Summary", level=1)
    doc.add_paragraph(summary)
    doc.add_heading("Technical Skills", level=1)
    doc.add_paragraph(", ".join(skills))
    doc.add_heading("Work Experience", level=1)
    for role, comp, dates, desc in experience:
        p = doc.add_paragraph()
        p.add_run(role).bold = True
        p.add_run(f" | {comp} ({dates})").italic = True
        doc.add_paragraph(desc)
    doc.add_heading("Education", level=1)
    doc.add_paragraph(education)
    doc.save(str(filepath))


def build_txt(filepath: Path, name: str, contact: str, summary: str, skills: list, experience: list, education: str):
    lines = [
        "=======================================================",
        name,
        contact,
        "=======================================================",
        "",
        "PROFESSIONAL SUMMARY:",
        summary,
        "",
        "TECHNICAL SKILLS:",
        ", ".join(skills),
        "",
        "WORK EXPERIENCE:"
    ]
    for role, comp, dates, desc in experience:
        lines.append(f"- {role} at {comp} ({dates}):")
        lines.append(f"  {desc}")
        lines.append("")
    lines.append("EDUCATION:")
    lines.append(education)
    lines.append("=======================================================")
    filepath.write_text("\n".join(lines), encoding="utf-8")


def generate_all():
    cands_file = DATA_DIR / "candidates.json"
    if not cands_file.exists():
        print(f"Error: {cands_file} not found.")
        return

    candidates = json.loads(cands_file.read_text(encoding="utf-8"))
    print(f"Generating {len(candidates)} resume files...")
    for c in candidates:
        fn = c["filename"]
        target = RESUMES_DIR / fn
        if fn.endswith(".pdf"):
            build_pdf(target, c["name"], c["contact"], c["summary"], c["skills"], c["experience"], c["education"])
        elif fn.endswith(".docx"):
            build_docx(target, c["name"], c["contact"], c["summary"], c["skills"], c["experience"], c["education"])
        elif fn.endswith(".txt"):
            build_txt(target, c["name"], c["contact"], c["summary"], c["skills"], c["experience"], c["education"])

    jds_file = DATA_DIR / "job_descriptions.json"
    if jds_file.exists():
        jds = json.loads(jds_file.read_text(encoding="utf-8"))
        for jd in jds:
            out_p = JDS_DIR / f"{jd['id']}.json"
            out_p.write_text(json.dumps(jd, indent=2), encoding="utf-8")

    print(f"Dataset successfully built in '{RESUMES_DIR}' and '{JDS_DIR}'.")


if __name__ == "__main__":
    generate_all()
